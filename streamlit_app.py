import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Pt
import tempfile
import os

# Configure the page
st.set_page_config(
    page_title="Sticker Label Generator",
    page_icon="🏷️",
    layout="wide"
)

st.title("🏷️ Sticker Label Generator")
st.markdown("Generate different types of labels from your CSV data")

# Dropdown for label type selection
label_type = st.selectbox(
    "Choose label type:",
    ["Side Labels", "Top Labels", "Tissue/Blood Tube Labels"],
    help="Select the type of labels you want to generate"
)

# File upload section
col1, col2 = st.columns(2)

with col1:
    st.subheader("📁 Upload CSV File")
    csv_file = st.file_uploader(
        "Choose a CSV file",
        type=['csv'],
        help="Upload your CSV data file"
    )

with col2:
    st.subheader("📋 Template Selection")
    template_option = st.radio(
        "Choose template source:",
        ["Use Built-in Template", "Upload Custom Template"],
        help="Use the built-in template or upload your own"
    )
    
    if template_option == "Upload Custom Template":
        template_file = st.file_uploader(
            "Choose a Word template file",
            type=['docx'],
            help="Upload your Word document template"
        )
    else:
        template_file = None
        st.info("Using built-in template for selected label type")

# Function to get template path based on label type
def get_template_path(label_type):
    template_mapping = {
        "Side Labels": "templates/side_label.docx",
        "Top Labels": "templates/top_label.docx", 
        "Tissue/Blood Tube Labels": "templates/tissue_label.docx"
    }
    return template_mapping.get(label_type)
def load_csv_with_encodings(csv_file):
    encodings = ['utf-8', 'ISO-8859-1', 'cp1252']
    
    for encoding in encodings:
        try:
            df = pd.read_csv(csv_file, encoding=encoding)
            st.success(f"Successfully read CSV with encoding: {encoding}")
            return df
        except Exception as e:
            st.warning(f"Failed to read with encoding {encoding}: {e}")
            continue
    
    st.error("Failed to read the CSV file with any encoding")
    return None

# Function for Side Labels formatting
def format_side_labels(row):
    return (
        f"{row['Species_code']} ({row['Sample Type']})\n"
        f"{row['Band_no']}\n"
        f"Ext: {row['Date']}\n"
        f"{row['CityTown']}, {row['State']} {row['Country_code']}"
    )

# Function for Top Labels formatting
def format_top_labels(row):
    # Split the BGP_ID into parts based on the delimiter 'N'
    bgp_id_parts = row['BGP_ID'].split('N')

    # Handle cases where 'N' is not found or there are more than two parts
    if len(bgp_id_parts) == 2:
        bgp_id_first_part = bgp_id_parts[0] + 'N'  # Add 'N' back to the first part
        bgp_id_second_part = bgp_id_parts[1]
    else:
        # If the delimiter 'N' is not found or there are more than 2 parts, handle accordingly
        bgp_id_first_part = row['BGP_ID']  # Keep it as-is if 'N' is not found
        bgp_id_second_part = ''

    # Return the formatted block of text
    return (
        f"{row['Species_code']}\n"
        f"{bgp_id_first_part}\n"
        f"{bgp_id_second_part}"
    )

# Function for Tissue/Blood Tube Labels formatting
def format_tissue_labels(row):
    # Split the BGP_ID into parts based on the delimiter 'N'
    bgp_id_parts = row['BGP_ID'].split('N')

    # Handle cases where 'N' is not found or there are more than two parts
    if len(bgp_id_parts) == 2:
        bgp_id_first_part = bgp_id_parts[0] + 'N'  # Add 'N' back to the first part
        bgp_id_second_part = bgp_id_parts[1]
    else:
        # If the delimiter 'N' is not found or there are more than 2 parts, handle accordingly
        bgp_id_first_part = row['BGP_ID']  # Keep it as-is if 'N' is not found
        bgp_id_second_part = ''

    # Return the formatted block of text
    return (
        f"\n"
        f"{bgp_id_first_part}\n"
        f"{bgp_id_second_part}"
    )

# Function to process labels based on type
def process_labels(df, doc, label_type):
    try:
        if not doc.tables:
            raise Exception("No tables found in the document.")
        
        table = doc.tables[0]
        
        # Get the appropriate formatting function
        if label_type == "Side Labels":
            format_function = format_side_labels
            output_filename = "side_labels_output.docx"
        elif label_type == "Top Labels":
            format_function = format_top_labels
            output_filename = "top_labels_output.docx"
        else:  # Tissue/Blood Tube Labels
            format_function = format_tissue_labels
            output_filename = "tissue_labels_output.docx"
        
        # Generate formatted texts
        formatted_texts = [format_function(row) for _, row in df.iterrows()]
        
        if label_type == "Tissue/Blood Tube Labels":
            # Special handling for tissue labels (every other row/column)
            num_rows = len(table.rows)
            num_cols = len(table.columns)
            
            # Check if we have enough space
            available_cells = (num_rows // 2 + (num_rows % 2)) * (num_cols // 2 + (num_cols % 2))
            if len(formatted_texts) > available_cells:
                st.warning(f"Not enough space in template. Template has {available_cells} available cells, but you have {len(formatted_texts)} records.")
            
            text_index = 0
            for row_idx in range(num_rows):
                if row_idx % 2 == 0:  # Skip rows to print only every other row
                    for col_idx in range(0, num_cols, 2):  # Print into every other cell
                        if text_index < len(formatted_texts):
                            cell = table.cell(row_idx, col_idx)
                            text = formatted_texts[text_index]
                            text_index += 1

                            # Clear existing text in the cell
                            for paragraph in cell.paragraphs:
                                if paragraph._element.getparent() is not None:
                                    cell._element.remove(paragraph._element)

                            # Add the formatted text to the cell
                            p = cell.add_paragraph()
                            run = p.add_run(text)
                            run.font.size = Pt(5.5)
        else:
            # Standard processing for Side and Top labels
            table_cells = [cell for row in table.rows for cell in row.cells]
            
            # Check if we have enough cells
            if len(formatted_texts) > len(table_cells):
                st.warning(f"Not enough cells in template. Template has {len(table_cells)} cells, but you have {len(formatted_texts)} records.")
            
            # Update each cell in the table
            for cell, text in zip(table_cells, formatted_texts):
                # Clear existing text in the cell
                for paragraph in cell.paragraphs:
                    if paragraph._element.getparent() is not None:
                        cell._element.remove(paragraph._element)

                # Add the formatted text to the cell
                p = cell.add_paragraph()
                run = p.add_run(text)
                run.font.size = Pt(5.5)
        
        return doc, output_filename
        
    except Exception as e:
        st.error(f"Error processing labels: {e}")
        return None, None

# Display CSV preview if uploaded
if csv_file is not None:
    df = load_csv_with_encodings(csv_file)
    
    if df is not None:
        # Clean column names
        df.columns = df.columns.str.strip()
        
        st.subheader("📊 CSV Data Preview")
        st.dataframe(df.head(10))
        st.info(f"Loaded {len(df)} rows and {len(df.columns)} columns")
        
        # Show required columns based on label type
        with st.expander("Required Columns for Selected Label Type"):
            if label_type == "Side Labels":
                required_cols = ['Species_code', 'Sample Type', 'Band_no', 'Date', 'CityTown', 'State', 'Country_code']
                st.write("Required columns:", required_cols)
                missing_cols = [col for col in required_cols if col not in df.columns]
                if missing_cols:
                    st.error(f"Missing required columns: {missing_cols}")
                else:
                    st.success("All required columns present!")
            
            elif label_type == "Top Labels":
                required_cols = ['Species_code', 'BGP_ID']
                st.write("Required columns:", required_cols)
                missing_cols = [col for col in required_cols if col not in df.columns]
                if missing_cols:
                    st.error(f"Missing required columns: {missing_cols}")
                else:
                    st.success("All required columns present!")
            
            else:  # Tissue/Blood Tube Labels
                required_cols = ['BGP_ID']
                st.write("Required columns:", required_cols)
                missing_cols = [col for col in required_cols if col not in df.columns]
                if missing_cols:
                    st.error(f"Missing required columns: {missing_cols}")
                else:
                    st.success("All required columns present!")
        
        # Show available columns
        with st.expander("Available Columns in Your CSV"):
            st.write(list(df.columns))

# Generate labels button
if csv_file is not None and df is not None:
    # Check if we have a template (either uploaded or built-in)
    has_template = template_file is not None or template_option == "Use Built-in Template"
    
    if has_template:
        if st.button("🚀 Generate Labels", type="primary"):
            with st.spinner(f"Generating {label_type.lower()}..."):
                try:
                    # Load the appropriate template
                    if template_option == "Use Built-in Template":
                        template_path = get_template_path(label_type)
                        if os.path.exists(template_path):
                            doc = Document(template_path)
                        else:
                            st.error(f"Built-in template not found: {template_path}")
                            st.stop()
                    else:
                        doc = Document(template_file)
                    
                    # Process the labels
                    processed_doc, output_filename = process_labels(df, doc, label_type)
                    
                    if processed_doc is not None:
                        # Save to temporary file
                        with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp_file:
                            processed_doc.save(tmp_file.name)
                            
                            # Read file for download
                            with open(tmp_file.name, 'rb') as f:
                                doc_bytes = f.read()
                            
                            # Clean up temp file
                            os.unlink(tmp_file.name)
                        
                        st.success(f"{label_type} generated successfully!")
                        
                        # Download button
                        st.download_button(
                            label="📥 Download Generated Labels",
                            data=doc_bytes,
                            file_name=output_filename,
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        )
                    
                except Exception as e:
                    st.error(f"Error generating labels: {str(e)}")
                    st.exception(e)
    else:
        st.warning("Please select a template option or upload a custom template.")
else:
    if csv_file is None:
        st.info("👆 Please upload a CSV file to get started.")

# Instructions
with st.expander("ℹ️ How to use this app"):
    st.markdown("""
    ### Instructions:
    1. **Select label type** from the dropdown (Side Labels, Top Labels, or Tissue/Blood Tube Labels)
    2. **Upload your CSV file** containing the data
    3. **Upload your Word template** (.docx file with the table layout)
    4. **Check the required columns** section to ensure your CSV has the needed data
    5. **Click "Generate Labels"** to process your data
    6. **Download the generated document** with your formatted labels
    
    ### Required Columns by Label Type:
    - **Side Labels**: Species_code, Sample Type, Band_no, Date, CityTown, State, Country_code
    - **Top Labels**: Species_code, BGP_ID
    - **Tissue/Blood Tube Labels**: BGP_ID
    
    ### Template Requirements:
    - Word document (.docx) with at least one table
    - The table should have enough cells for your data
    - For Tissue/Blood Tube Labels, the template uses every other row and column
    """)

# Footer
st.markdown("---")
st.markdown("🏷️ Sticker Label Generator - Built with Streamlit")